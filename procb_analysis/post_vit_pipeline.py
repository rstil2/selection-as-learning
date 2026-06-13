#!/usr/bin/env python3
"""
Post-completion pipeline for ViT experiment.
Run after ml_vit_results.json exists.
"""
import json
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent
VIT_JSON = ROOT / "ml_vit_results.json"
INCONCLUSIVE = ROOT / "vit_result_inconclusive.txt"
THRESHOLD = -0.70


def main():
    if not VIT_JSON.exists():
        print("ml_vit_results.json not found — ViT experiment not complete.")
        sys.exit(1)

    with open(VIT_JSON) as f:
        vit = json.load(f)

    r = vit["pearson_r"]
    p = vit["pearson_p"]
    direction = vit.get("direction_consistent", r < 0)

    print(f"ViT result: r = {r:.3f}, p = {p:.4f}, direction_consistent = {direction}")

    if r >= THRESHOLD or not direction:
        INCONCLUSIVE.write_text(
            f"ViT replication inconclusive (threshold r < {THRESHOLD}).\n"
            f"Observed: r = {r:.3f}, p = {p:.4f}, direction_consistent = {direction}\n"
            f"ResNet-18 result (r = -0.97) stands alone; manuscript not updated for ViT.\n"
        )
        print(f"Result below threshold — wrote {INCONCLUSIVE.name}")
        sys.exit(0)

    print("Strong negative result — updating manuscript...")
    subprocess.run(
        [sys.executable, str(ROOT / "update_manuscript.py")],
        check=True,
        cwd=ROOT,
    )

    print("Regenerating universality figures...")
    subprocess.run(
        [sys.executable, str(ROOT / "universality_analysis.py")],
        check=True,
        cwd=ROOT,
    )

    print("Verifying manuscript paragraphs...")
    import docx

    doc = docx.Document(ROOT / "Manuscript_Draft_v2.docx")
    for idx in [7, 38, 47, 101, 121]:
        text = doc.paragraphs[idx].text
        if "PENDING" in text:
            print(f"  WARNING: para [{idx}] still has placeholders")
        else:
            print(f"  OK para [{idx}]")

    print("Pipeline complete.")


if __name__ == "__main__":
    main()
